from pandas import read_excel
from docx import Document
import tkinter as tk
from tkinter import filedialog, messagebox

def word_table(excel_file: str, n_rows: int, n_cols: int, file_name: str):
    df = read_excel(excel_file, nrows=n_rows).astype(str)
    df = df.iloc[:, :n_cols]
    pdf = Document()
    table = pdf.add_table(rows=len(df.index) + 1, cols=len(df.columns))  # نزيد rows 1 لكي نضيف خانه لأسماء الأعمدة

    for j, col_name in enumerate(df.columns):  # يمر j على ارقام الأعمدة و col يمر على اسماء الأعمدة
        table.cell(0, j).text = col_name  # cell هي الخلية التي يضع فيها اسم الأعمدة و تكون 0 لأنها الصف الاول في الجدول

    for i in range(df.shape[0]):  # يمر على الصفوف
        for j in range(df.shape[1]):  # يمر على الأعمدة
            table.cell(i + 1, j).text = repr(df.iat[i, j]).replace("'", "")  # نزيد الخلية 1 لأن الخلية 0 للأعمدة

    pdf.save(f"{file_name}.docx")

def browse_file():
    file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xls;*.xlsx")])
    excel_file_entry.delete(0, tk.END)
    excel_file_entry.insert(0, file_path)

def browse_file_1():
    file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xls;*.xlsx")])
    excel_file_entry.delete(0, tk.END)
    excel_file_entry.insert(0, file_path)

def convert_file():
    try:
        excel_file = excel_file_entry.get()
        n_rows = int(rows_entry.get())
        n_cols = int(cols_entry.get())
        file_name = new_file_entry.get()

        word_table(excel_file, n_rows, n_cols, file_name)
        messagebox.showinfo("نجاح", f"تم إنشاء الملف بنجاح: {file_name}.docx")
    except ValueError:
        messagebox.showerror("خطأ", "تأكد من إدخال عدد صحيح للصفوف والأعمدة.")
    except FileNotFoundError:
        messagebox.showerror("خطأ", "تأكد من أن ملف Excel موجود ويمكن الوصول إليه.")
    except Exception as e:
        messagebox.showerror("خطأ غير متوقع", str(e))

def close_app():
    root.destroy()

def main():
    global root, excel_file_entry, rows_entry, cols_entry, new_file_entry

    # إنشاء نافذة التطبيق
    root = tk.Tk()
    root.title("تحويل ملف Excel إلى جدول Word")

    # إنشاء وإضافة عناصر الواجهة
    tk.Label(root, text="ملف Excel:").grid(row=0, column=0, sticky=tk.W, pady=2)
    excel_file_entry = tk.Entry(root, width=40)
    excel_file_entry.grid(row=0, column=1, pady=2)
    tk.Button(root, text="استعراض", command=browse_file).grid(row=0, column=2, pady=2)

    tk.Label(root, text="عدد الصفوف:").grid(row=1, column=0, sticky=tk.W, pady=2)
    rows_entry = tk.Entry(root, width=10)
    rows_entry.grid(row=1, column=1, pady=2)

    tk.Label(root, text="عدد الأعمدة:").grid(row=2, column=0, sticky=tk.W, pady=2)
    cols_entry = tk.Entry(root, width=10)
    cols_entry.grid(row=2, column=1, pady=2)

    tk.Label(root, text="اسم الملف الجديد:").grid(row=3, column=0, sticky=tk.W, pady=2)
    new_file_entry = tk.Entry(root, width=20)
    new_file_entry.grid(row=3, column=1, pady=2)

    tk.Label(root, text=" -- يجب تعبئة جميع الحقول --").grid(row=4, columnspan=3, pady=2)

    # وضع الأزرار في أسفل النافذة في المنتصف
    button_frame = tk.Frame(root)
    button_frame.grid(row=5, columnspan=3, pady=10)

    tk.Button(button_frame, text="تأكيد", command=convert_file, width=10).grid(row=0, column=0, padx=5)
    tk.Button(button_frame, text="إلغاء", command=close_app, width=10).grid(row=0, column=1, padx=5)

    # تشغيل التطبيق
    root.mainloop()

if __name__ == "__main__":
    main()
